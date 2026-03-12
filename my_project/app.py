from flask import Flask, render_template, request, jsonify, send_file, abort
import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
import openpyxl.styles
import openpyxl.utils
from io import BytesIO
from datetime import datetime
from services.api_service import APIService

app = Flask(__name__)
api_service = APIService()

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/photo/<path:filename>')
def get_photo(filename):
    try:
        if os.path.exists(filename):
            return send_file(filename)
        else:
            return abort(404)
    except Exception:
        return abort(404)

@app.route('/api/load-data', methods=['GET'])
def load_local_data():
    try:
        config_data = api_service.concurrent_service.load_data_config()
        return jsonify({
            'success': True,
            'data': config_data.get('data', []),
            'styles': config_data.get('style', ['文艺']),
            'default_style': config_data.get('default_style', '文艺'),
            'system_prompt': config_data.get('system_prompt', '')
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/concurrent', methods=['POST'])
def send_concurrent_requests():
    try:
        data = request.get_json() or {}
        user_prompt = data.get('userPrompt', '')
        style = data.get('style', '文艺')
        system_prompt = data.get('systemPrompt', '')
        results = api_service.send_concurrent_requests(
            user_prompt=user_prompt,
            style=style,
            system_prompt=system_prompt
        )
        return jsonify({
            'success': True,
            'total': len(results),
            'results': results
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    try:
        data = request.get_json() or {}
        results = data.get('results', [])
        
        wb = Workbook()
        ws = wb.active
        ws.title = '文案对比结果'
        
        headers = ['数据项ID', '状态', '图片1', '图片2', '图片3', '图片4', '图片5', '原始文案', '返回文案']
        ws.append(headers)
        
        for col_idx in range(2, 7):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx+1)].width = 15
        
        for idx, item in enumerate(results, 1):
            data_id = item.get('data_id', idx)
            success = item.get('success', False)
            status = '成功' if success else '失败'
            
            photos = item.get('result', {}).get('photos', [])
            photo_urls = [p.get('url', '') for p in photos]
            
            while len(photo_urls) < 5:
                photo_urls.append('')
            
            original_text = item.get('result', {}).get('text', '')
            
            if success:
                generated_text = item.get('result', {}).get('generated_text', '')
            else:
                generated_text = '错误：' + item.get('error', '未知错误')
            
            row = [data_id, status, '', '', '', '', '', original_text, generated_text]
            ws.append(row)
            
            ws.row_dimensions[idx+1].height = 100
            
            for col_idx, photo_url in enumerate(photo_urls, 1):
                if photo_url and os.path.exists(photo_url):
                    try:
                        img = ExcelImage(photo_url)
                        img.width = 80
                        img.height = 80
                        cell_letter = openpyxl.utils.get_column_letter(col_idx+2)
                        cell_coord = f"{cell_letter}{idx+1}"
                        ws.add_image(img, cell_coord)
                        import sys
                        print(f"Added image to {cell_coord}: {photo_url}", file=sys.stderr)
                    except Exception as e:
                        import sys
                        print(f"Error loading image {photo_url}: {e}", file=sys.stderr)
                        pass
        
        import sys
        print(f"Total images added to worksheet: {len(ws._images)}", file=sys.stderr)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"copywriting_result_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        
        from flask import Response
        from urllib.parse import quote
        return Response(
            output.read(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{quote(filename)}"'}
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
